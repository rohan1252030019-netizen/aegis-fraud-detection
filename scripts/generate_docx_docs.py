"""
AEGIS - Word (.docx) Documentation Generator
Builds a professional, comprehensive User Interface & System Architecture guide.
"""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_document():
    doc = Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color palette
    PRIMARY_BLUE = RGBColor(30, 64, 175)     # #1e40af
    DARK_TEXT = RGBColor(30, 41, 59)          # #1e293b
    MUTED_TEXT = RGBColor(100, 116, 139)      # #64748b

    def heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(17)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_BLUE
        return p

    def heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_BLUE
        return p

    def add_p(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK_TEXT
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK_TEXT
        return p

    def add_callout(text, label="NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        r1 = p.add_run(f"{label}: ")
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_BLUE
        r2 = p.add_run(text)
        r2.font.size = Pt(10)
        r2.font.color.rgb = DARK_TEXT
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def format_table(tbl, col_widths, headers, data):
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Headers
        hdr_cells = tbl.rows[0].cells
        for i, title in enumerate(headers):
            hdr_cells[i].text = title
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            hdr_cells[i].paragraphs[0].paragraph_format.space_before = Pt(4)
            hdr_cells[i].paragraphs[0].paragraph_format.space_after = Pt(4)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
            hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        
        # Rows
        for r_idx, row_data in enumerate(data):
            row = tbl.add_row()
            fill_color = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(val)
                p = cell.paragraphs[0]
                if len(p.runs) > 0:
                    p.runs[0].font.size = Pt(9.5)
                    p.runs[0].font.color.rgb = DARK_TEXT
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
                cell._tc.get_or_add_tcPr().append(shd)
                
        # Set widths
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # Document Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(15)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run("AEGIS: Financial Fraud & Money Mule Detection System")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = PRIMARY_BLUE

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(8)
    r_sub = sub_p.add_run("Complete User Interface Guide, Dashboard Telemetry, and Architecture Manual")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = MUTED_TEXT

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    r_meta = meta_p.add_run("Live Website: https://aegis-edi.vercel.app\nDepartment of CSE (AI & ML), VIT Pune | EDI Project")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(10)
    r_meta.font.italic = True
    r_meta.font.color.rgb = PRIMARY_BLUE

    add_callout(
        "This document is prepared to provide evaluators, faculty, and compliance analysts with an easy-to-understand explanation of every card, metric, button, and visual widget across the AEGIS platform.",
        "EXECUTIVE SUMMARY",
    )

    # Section 1
    heading_1("1. Project Overview & Biomimetic Concept")
    add_p(
        "Money laundering syndicates deploy mule accounts to funnel illicit funds through automated micro-transfers (smurfing) and rapid cash-out conduits. Traditional banking systems rely on static thresholds or isolated machine-learning classifiers that produce high false-positive rates and zero explainability."
    )
    add_p(
        "AEGIS implements a biomimetic multi-layer defense framework modeled on the human immune system: Innate Immunity (rapid temporal velocity rules), Adaptive Immunity (sequence Transformer deep learning, unsupervised Isolation Forests, and NetworkX topological graph analysis), and Formal Immune Memory (vector-based historical threat retrieval, mathematical Z3 invariant proofs, and on-premise local LLM intelligence briefs)."
    )

    # Section 2: Dashboard
    heading_1("2. Screen 1: Security & Fraud Overview Dashboard (/dashboard)")
    add_p(
        "The Dashboard is the primary operational console giving compliance teams real-time visibility into overall financial health, anomalous velocity, and active risk."
    )

    heading_2("A. Top Primary Stat Cards")
    stat_cols = [1.8, 1.2, 3.8]
    stat_headers = ["Stat Card", "Sample Value", "What It Represents & Real-World Impact"]
    stat_data = [
        [
            "Monitored Accounts",
            "1,959",
            "Total pool of customer and corporate accounts actively monitored. Includes 30-day growth trend (+3.8%).",
        ],
        [
            "Mule Candidates",
            "14",
            "High-priority accounts with composite threat scores >= 70/100, flagged for rapid fund dispersal or pass-through behavior.",
        ],
        [
            "Active Alerts",
            "23",
            "Unresolved anomalies detected across Temporal, Behavioral, and Graph layers awaiting compliance triage.",
        ],
        [
            "Pending Dossiers",
            "8",
            "Formal investigation cases compiled and currently under review by human investigators before filing a SAR.",
        ],
    ]
    stat_tbl = doc.add_table(rows=1, cols=3)
    format_table(stat_tbl, stat_cols, stat_headers, stat_data)

    heading_2("B. Analytics Charts & Widgets")
    add_p(
        "1. Transaction Velocity & Anomaly Trends (Area Chart): Plots total 24-hour transaction flow (Cyan area) alongside flagged suspicious transaction volume (Rose area). Sudden spikes in the rose area reveal automated fund extraction batches, often occurring during off-peak hours."
    )
    add_p(
        "2. Risk Distribution by Tier (Bar Chart): Visualizes account population across 4 calibrated tiers: Low Risk (Green, benign users), Moderate (Yellow, slight deviation), Elevated (Orange, nearing velocity barriers), and Critical/High (Red, verified mule candidates)."
    )
    add_p(
        "3. Biomimetic Defense System Status Card: Displays active health checks for Innate Immunity (Rules), Adaptive Immunity (Transformers), Network Topology (Graph), and Formal Grounding (Z3 Proofs)."
    )
    add_p(
        "4. Sync Engine Button: Located in the top header. Clicking this forces a live synchronization with the detection engine to refresh metrics."
    )

    # Section 3: Data Ingestion
    heading_1("3. Screen 2: Data Ingestion & Real-Time Validation (/data-import)")
    add_p(
        "The Data Ingestion module enables users to upload raw financial transaction batches (.CSV or .JSON) and execute multi-layer detection pipelines."
    )

    heading_2("A. Pre-Validation Safeguards")
    add_p(
        "• Strict 250 MB Size Limit: Enforces a strict ceiling on uploaded files, preventing memory exhaustion and reporting exact byte counts."
    )
    add_p(
        "• Instant Schema Inspector: The moment a file is dropped into the zone, AEGIS inspects its header to verify the 6 Canonical Fields: transaction_id, sender_account_id, receiver_account_id, amount, currency, and timestamp."
    )
    add_p(
        "• Schema Badges: Displays 'Ready for Ingestion (AEGIS_CANONICAL)' in green when valid, or clearly lists missing fields when incompatible schemas are detected."
    )

    heading_2("B. Execution Results & Quality Scoring")
    add_p(
        "• Run AEGIS Detection Pipeline Button: Triggers the ingestion engine with an animated progress bar from 0% to 100%."
    )
    add_p(
        "• Dataset Validation Card: Displays total processed records (e.g. 663,373), valid count, rejected count, Quality Score (100.0% EXCELLENT), and an expandable row-by-row error inspector."
    )
    add_p(
        "• Pipeline Summary Counters: Confirms transactions processed, accounts analyzed, anomalies detected, suspicious networks, alerts raised, and cases generated."
    )

    # Section 4: Alerts & Accounts
    heading_1("4. Screens 3, 4 & 5: Alerts, Monitored Accounts, and Transactions")
    add_p(
        "• Active Alerts Inbox (/alerts): Features severity tabs (ALL, CRITICAL, HIGH, MEDIUM, LOW) and account search. Each alert displays a unique ID with one-click copy, threat type (e.g. MULTI_LAYER_MULE_RISK, STRUCTURING_SMURFING, CYCLE_TOPOLOGY), risk score, and status."
    )
    add_p(
        "• Monitored Accounts Directory (/accounts): Lists customer accounts with their composite risk scores, status (FLAGGED, UNDER_REVIEW, FROZEN, ACTIVE), and total transaction volume (Sent/Received). Clicking an account navigates directly to its topological graph."
    )
    add_p(
        "• Transaction Ledger (/transactions): Live audit ledger displaying sender, receiver, amount, currency, ISO timestamp, and flag status. Includes a 'Flagged Only' toggle for rapid auditing."
    )

    # Section 5: Case Investigation
    heading_1("5. Screen 6: Case Management & 12-Stage Investigation Engine (/cases)")
    add_p(
        "The Investigation Workbench is the flagship module of AEGIS, providing full compliance transparency and explainability."
    )

    heading_2("A. The 12 Autonomous Pipeline Stages")
    stages = (
        "1. Data Ingestion: Streams transaction records.\n"
        "2. Preprocessing & Data Quality: Cleans missing values and standardizes datetimes.\n"
        "3. Temporal Velocity & Structuring: Flags transfers under statutory reporting thresholds.\n"
        "4. Transformer Sequence Inference: PyTorch sequence anomaly scoring.\n"
        "5. Behavioral Outlier Detection: Scikit-learn Isolation Forest profiling 8 dimensions.\n"
        "6. Graph Correlation & Cycle Analysis: NetworkX cycle and conduit detection.\n"
        "7. Multi-Layer Evidence Fusion: Synthesizes scores into calibrated 0-100 composite risk.\n"
        "8. SHAP Feature Attribution: TreeExplainer computing exact risk driver contributions.\n"
        "9. Formal Logic Verification: Mathematical Z3 invariant theorem grounding.\n"
        "10. Threat Memory Retrieval: Vector cosine similarity matching against known syndicates.\n"
        "11. Local Investigation Intelligence: Local Ollama / Llama 3 generating narrative briefs.\n"
        "12. Final Compliance Dossier Compilation: Assembles regulatory SAR artifact."
    )
    add_p(stages)

    heading_2("B. Dossier Components & Formal Grounding")
    add_p(
        "• Risk Profile Card: Displays composite risk score (e.g. 92.4/100), classification ('Coordinated Smurfing & Mule Aggregator'), and confidence (94%)."
    )
    add_p(
        "• Formal Verification Proofs (Z3): Proves mathematically that AML invariants hold true with zero counterexamples: RapidDisbursement (funds drained <300s), StructuringThresholdBarrier, and CircularTopology."
    )
    add_p(
        "• SHAP Drivers: Tells investigators why the account was flagged (+0.38 fan_in_velocity, +0.29 structuring_proximity)."
    )
    add_p(
        "• Threat Memory Corroboration: Shows 93% cosine match with historical Cryptocurrency Mule Syndicate 081."
    )
    add_p(
        "• Local LLM Brief: Formats a complete regulatory narrative without leaking banking data to cloud APIs."
    )
    add_p(
        "• Regulatory Actions: Buttons for 'Confirm Mule (True Positive)' and 'Mark False Positive', updating threat memory in real-time."
    )

    # Section 6: Graph Visualization
    heading_1("6. Screen 7: Interactive Financial Network Graph (/graph)")
    add_p(
        "Powered by dynamic D3.js force-directed simulations, this module visualizes complex laundering topologies:"
    )
    add_p(
        "• Red Nodes (Mule Aggregators): High-centrality accounts accumulating multiple inbound wires."
    )
    add_p(
        "• Amber Nodes (Smurfing Sources): Smurf accounts sending structured amounts below legal limits."
    )
    add_p(
        "• Outlined Gateways: Offramp crypto merchants or high-risk OTC brokers where laundered funds exit."
    )
    add_p(
        "• Interactive Physics: Users can zoom, pan, drag nodes, and click any account to view degree, betweenness centrality, and risk score."
    )

    # Section 7: User Workflows
    heading_1("7. Step-by-Step User Workflows")
    add_p(
        "Workflow 1: Ingesting a Dataset\n1. Go to /data-import -> Drag transaction CSV -> Verify 'Ready for Ingestion' badge -> Click 'Run AEGIS Detection Pipeline' -> View validation report."
    )
    add_p(
        "Workflow 2: Investigating an Account\n1. Go to /cases -> Select a case (e.g. CASE-20260914-8599) -> Click 'Run Full AEGIS Investigation' -> Watch 12 stages run -> Review Z3 theorems & SHAP drivers -> Click 'Confirm Mule' to record regulatory decision."
    )

    # Save
    out_dir = r"c:\Users\ADMIN\Documents\EDI PROJECT\aegis\docs"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "AEGIS_Project_User_Guide_and_Documentation.docx")
    doc.save(out_file)
    print(f"Successfully generated DOCX at: {out_file}")

if __name__ == "__main__":
    create_document()
